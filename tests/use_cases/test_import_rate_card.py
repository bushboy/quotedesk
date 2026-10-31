import csv
from decimal import Decimal
from io import BytesIO, StringIO
from uuid import UUID

import pytest
from openpyxl import Workbook

from quotedesk.adapters.outbound.parsers.macros import MacroRejected, reject_macros_before_parse
from quotedesk.application.dto.published import PublishedPolicy
from quotedesk.application.dto.tenant import TenantContext
from quotedesk.application.use_cases.imports.import_rate_card import ImportRateCard, ImportRateCardInput
from quotedesk.domain.policy.clauses import ClausePolicy
from quotedesk.domain.policy.draft import DraftRateRow, PolicyDraft
from quotedesk.domain.services.price import RateCard, RateCardItem
from quotedesk.domain.shared.errors import ValidationError
from quotedesk.domain.shared.money import Money
from tests.unit.domain.water_policy import aquasure_policy

TENANT = UUID("11111111-1111-4111-8111-111111111111")
CONTEXT = TenantContext(tenant_id=TENANT)
SENTENCE = "Call us about unusual shapes."
MACRO = "Macro-enabled files aren't accepted. Save a copy without macros and send that."
HEADER = ["item_code", "unit_price", "unit"]
CARD = [
    HEADER,
    ["TANK-5000", 4250.50, "each"],
    ["PUMP-STD", 1800.00, "each"],
]


class Mislabel:
    def __init__(self) -> None:
        self.headers: tuple[str, ...] | None = None

    def label(self, headers) -> dict[str, str]:
        self.headers = tuple(headers)
        return {"unit_price": "item_code"}


class Library:
    def __init__(self) -> None:
        self.live = PublishedPolicy(
            policy=aquasure_policy(),
            rate_card=RateCard(
                vat_inclusive=True,
                site_visit_fee=Money("450.00"),
                items=(RateCardItem(item_code="TANK-5000", description="5000 L tank", unit_price=Money("999.00")),),
            ),
            clause_policy=ClausePolicy(
                templates=(),
                settings={"deposit_percent": 30, "validity_days": 14},
                rate_card_codes=("TANK-5000",),
            ),
        )
        self.current_draft = PolicyDraft(
            deposit_percent=30,
            discount_cap_percent=5,
            validity_days=14,
            vat_inclusive=True,
            rate_rows=(DraftRateRow(item_code="OLD-ITEM", unit_price=Decimal("1.00"), vat_inclusive=True),),
            clauses=(),
            required_facts=(),
            rules=(),
            settings=("deposit_percent", "validity_days"),
        )
        self.saves = 0

    def published(self, _tenant_id: UUID) -> PublishedPolicy:
        return self.live

    def draft(self, _tenant_id: UUID) -> PolicyDraft:
        return self.current_draft

    def save_draft(self, _tenant_id: UUID, draft: PolicyDraft) -> None:
        self.saves += 1
        self.current_draft = draft


class Unit:
    def __init__(self) -> None:
        self.commits = 0
        self.rollbacks = 0

    def commit(self) -> None:
        self.commits += 1

    def rollback(self) -> None:
        self.rollbacks += 1


def _world() -> tuple[Library, Unit, ImportRateCard]:
    library = Library()
    unit = Unit()
    return library, unit, ImportRateCard(library, unit, reject_macros_before_parse)


def _xlsx(rows: list[list[object]]) -> bytes:
    book = Workbook()
    sheet = book.active
    for row_number, row in enumerate(rows, start=1):
        for column, value in enumerate(row, start=1):
            sheet.cell(row_number, column, value)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _csv(rows: list[list[object]]) -> bytes:
    buffer = StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows([["" if value is None else value for value in row] for row in rows])
    return buffer.getvalue().encode()


def _run(library_unit: tuple[Library, Unit, ImportRateCard] | None = None, **fields):
    library, unit, importer = _world() if library_unit is None else library_unit
    payload: dict[str, object] = {"expected_unit": "each", "labeller": Mislabel()}
    payload.update(fields)
    data = ImportRateCardInput(**payload)
    return importer.execute(CONTEXT, data), library, unit


def test_a_sheet_without_applies_to_stores_both() -> None:
    library, _unit, importer = _world()
    importer.execute(
        CONTEXT,
        ImportRateCardInput(filename="card.xlsx", body=_xlsx(CARD), expected_unit="each", labeller=Mislabel()),
    )
    assert [row.applies_to for row in library.current_draft.rate_rows] == ["both", "both"]

    stated = [
        ["item_code", "unit_price", "unit", "applies_to"],
        ["STAND", 900, "each", "new"],
    ]
    library, _unit, importer = _world()
    importer.execute(
        CONTEXT,
        ImportRateCardInput(filename="stand.xlsx", body=_xlsx(stated), expected_unit="each", labeller=Mislabel()),
    )
    assert library.current_draft.rate_rows[0].applies_to == "new"


def test_excel_numbers_come_from_cells() -> None:
    labeller = Mislabel()
    library, _unit, importer = _world()
    result = importer.execute(
        CONTEXT,
        ImportRateCardInput(filename="card.xlsx", body=_xlsx(CARD), expected_unit="each", labeller=labeller),
    )

    assert labeller.headers is not None
    assert "unit_price" in labeller.headers
    assert labeller.label(labeller.headers) == {"unit_price": "item_code"}
    assert [(row.item_code, row.unit_price, row.cell, row.status) for row in result.rows] == [
        ("TANK-5000", Decimal("4250.50"), "B2", "ready"),
        ("PUMP-STD", Decimal("1800.00"), "B3", "ready"),
    ]
    assert [row.unit_price for row in library.current_draft.rate_rows] == [Decimal("4250.50"), Decimal("1800.00")]


def test_macro_files_rejected(monkeypatch: pytest.MonkeyPatch) -> None:
    called: list[str] = []

    def refuse_xlsx(*_args, **_kwargs):
        called.append("xlsx")
        raise AssertionError("parsed")

    def refuse_csv(*_args, **_kwargs):
        called.append("csv")
        raise AssertionError("parsed")

    monkeypatch.setattr("openpyxl.load_workbook", refuse_xlsx)
    monkeypatch.setattr(csv, "reader", refuse_csv)
    for name in ("card.xlsm", "card.docm", "card.pptm", "quotes/card.XLSM"):
        library, _unit, importer = _world()
        with pytest.raises(MacroRejected) as caught:
            importer.execute(CONTEXT, ImportRateCardInput(filename=name, body=b"not-empty", expected_unit="each", labeller=Mislabel()))
        assert caught.value.code == "import.macro"
        assert caught.value.sentence == MACRO
        assert library.saves == 0
        assert library.current_draft.rate_rows[0].item_code == "OLD-ITEM"
    for name in ("card.pdf", "card.docx", "card.pptx", "card.png", "card.jpg"):
        library, _unit, importer = _world()
        with pytest.raises(ValidationError, match="That file type isn't imported yet.") as caught:
            importer.execute(CONTEXT, ImportRateCardInput(filename=name, body=b"not-empty", expected_unit="each", labeller=Mislabel()))
        assert caught.value.code == "import.unsupported"
        assert library.saves == 0
    assert called == []


def test_csv_and_xlsx_same_rows() -> None:
    csv_rows = [HEADER, ["TANK-5000", "4250.50", "each"], ["PUMP-STD", "1800.00", "each"]]
    from_xlsx, _, _ = _run(filename="card.xlsx", body=_xlsx(CARD))
    from_csv, _, _ = _run(filename="card.csv", body=_csv(csv_rows))

    assert [(row.item_code, row.unit_price, row.unit, row.flags, row.status, row.cell) for row in from_csv.rows] == [
        (row.item_code, row.unit_price, row.unit, row.flags, row.status, row.cell) for row in from_xlsx.rows
    ]


def test_import_flags_incomplete_rows() -> None:
    missing_and_blank = [
        HEADER,
        ["TANK-5000", None, "each"],
        ["PUMP-STD", "1800.00", None],
    ]
    result, _, _ = _run(filename="card.xlsx", body=_xlsx(missing_and_blank))
    missing, blank = result.rows
    assert missing.unit_price is None
    assert missing.flags == ("missing_price",)
    assert missing.status == "needs_review"
    assert blank.unit == ""
    assert blank.flags == ("blank_unit",)
    assert blank.unit_price == Decimal("1800.00")
    assert blank.status == "needs_review"

    mismatch, library, _ = _run(
        filename="card.csv",
        body=_csv([HEADER, ["WIN-STD", "850", "each"]]),
        expected_unit="m2",
    )
    row = mismatch.rows[0]
    assert row.flags == ("unit_mismatch",)
    assert row.status == "needs_review"
    assert row.unit_price == Decimal("850.00")
    assert library.current_draft.rate_rows[0].item_code == "WIN-STD"


def test_import_keeps_unmapped_text() -> None:
    rows = CARD + [[SENTENCE]]
    result, _, _ = _run(filename="card.xlsx", body=_xlsx(rows))

    assert result.unmapped_text == (SENTENCE,)
    assert SENTENCE not in {row.item_code for row in result.rows}
    assert [row.item_code for row in result.rows] == ["TANK-5000", "PUMP-STD"]


def test_import_conflict_is_not_averaged() -> None:
    rows = [HEADER, ["WIN-STD", 850, "each"], ["WIN-STD", 900, "each"]]
    result, library, _ = _run(filename="card.xlsx", body=_xlsx(rows))

    assert [row.unit_price for row in result.rows] == [Decimal("850.00"), Decimal("900.00")]
    assert [row.status for row in result.rows] == ["needs_review", "needs_review"]
    assert [row.flags for row in result.rows] == [("conflict",), ("conflict",)]
    assert Decimal("875") not in {row.unit_price for row in result.rows}
    assert [row.unit_price for row in library.current_draft.rate_rows] == [Decimal("850.00"), Decimal("900.00")]


def test_import_does_not_publish() -> None:
    library, unit, importer = _world()
    before = library.published(TENANT)
    result = importer.execute(
        CONTEXT,
        ImportRateCardInput(filename="card.xlsx", body=_xlsx(CARD), expected_unit="each", labeller=Mislabel()),
    )

    assert [row.item_code for row in result.rows] == ["TANK-5000", "PUMP-STD"]
    draft = library.draft(TENANT)
    assert [(row.item_code, row.unit_price) for row in draft.rate_rows] == [
        ("TANK-5000", Decimal("4250.50")),
        ("PUMP-STD", Decimal("1800.00")),
    ]
    assert draft.deposit_percent == 30
    published = library.published(TENANT)
    assert published.policy.version == before.policy.version
    assert published.rate_card.items[0].unit_price == Money("999.00")
    assert published is before
    assert unit.commits == 1


def test_oversize_and_empty_upload() -> None:
    library, _unit, importer = _world()
    with pytest.raises(ValidationError, match="That file is empty.") as empty:
        importer.execute(CONTEXT, ImportRateCardInput(filename="card.csv", body=b"", expected_unit="each", labeller=Mislabel()))
    assert empty.value.code == "import.empty"
    assert library.saves == 0
    assert library.current_draft.rate_rows[0].item_code == "OLD-ITEM"

    library, _unit, importer = _world()
    with pytest.raises(ValidationError, match="That file is too large.") as oversize:
        importer.execute(
            CONTEXT,
            ImportRateCardInput(filename="card.xlsx", body=b"x" * 2_000_001, expected_unit="each", labeller=Mislabel()),
        )
    assert oversize.value.code == "import.oversize"
    assert library.saves == 0
    assert library.current_draft.rate_rows[0].unit_price == Decimal("1.00")

    core = _csv([HEADER, ["TANK-5000", "4250.50", "each"]])
    body = core + (b"\n" * (2_000_000 - len(core)))
    assert len(body) == 2_000_000
    result, library, _unit = _run(filename="card.csv", body=body)
    assert result.rows[0].unit_price == Decimal("4250.50")
    assert library.current_draft.rate_rows[0].item_code == "TANK-5000"
